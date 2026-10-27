"""docx rendering for the JD pipeline (txt -> docx).

Moved out of jd_scorecard_resume_v2.py on 27 Oct 2026 so it can be imported and
run on its own (v2 needs an API key just to start, and must not be imported).

Why this exists / what changed:
  * Scorecards are written by the LLM as full Markdown (# headings, **bold**,
    tables, bullets, numbered lists, --- rules). The old writer treated every
    .txt as plain text, so literal "**", "##" and "---" ended up in the docx
    and therefore in the PDF made from it. ScoreCards are now rendered with a
    real Markdown renderer (render_markdown).
  * Resumes / cover letters keep the original layout code path, but no literal
    "**" can reach the document any more (headings, "Company — Title" lines
    and unbalanced markers are all cleaned).

Standalone use (rebuild docx from existing .txt, e.g. to clean old outputs):
    python scripts/docx_render.py <txt file | folder> [--no-backup]
The previous docx is copied to archive/docx_rebuild_<YYYYMMDD>/ first.
"""

import re
import shutil
import sys
from datetime import datetime
from pathlib import Path

try:
    from docx import Document
    from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

# **bold**  |  `code`  |  *italic*  (italic must hug its text and not touch word chars)
_INLINE_RE = re.compile(
    r"\*\*(.+?)\*\*|`([^`]+)`|(?<![\w*])\*(?![\s*])(.+?)(?<![\s*])\*(?![\w*])"
)


def _emit_run(paragraph, text, bold=False, italic=False, size=None, mono=False):
    text = text.replace("**", "")          # never let a marker reach the output
    if not text:
        return
    run = paragraph.add_run(text)
    if bold:
        run.bold = True
    if italic:
        run.italic = True
    if size:
        run.font.size = Pt(size)
    if mono:
        run.font.name = "Consolas"


def add_runs_with_markup(paragraph, text, bold=False, size=None):
    """Add text to a paragraph, rendering **bold**, *italic* and `code` spans as
    formatted runs. Unbalanced / leftover markers are removed, never printed."""
    pos = 0
    for m in _INLINE_RE.finditer(text):
        if m.start() > pos:
            _emit_run(paragraph, text[pos:m.start()], bold=bold, size=size)
        if m.group(1) is not None:
            _emit_run(paragraph, m.group(1), bold=True, size=size)
        elif m.group(2) is not None:
            _emit_run(paragraph, m.group(2), bold=bold, size=size, mono=True)
        else:
            _emit_run(paragraph, m.group(3), bold=bold, italic=True, size=size)
        pos = m.end()
    if pos < len(text):
        _emit_run(paragraph, text[pos:], bold=bold, size=size)


def style_docx_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    # python-docx's default template applies a document-wide 10pt "space after"
    # (w:docDefaults/w:pPrDefault) to every paragraph that doesn't explicitly
    # override it — including every blank spacer paragraph, section-title
    # line, and bullet. That silently stacked extra air on top of the blank
    # rows this converter deliberately inserts. Zero it at the Normal-style
    # level (List Bullet is based on Normal) so only the explicit Pt(3)/Pt(4)
    # overrides below ever add spacing.
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)


# ── Resume / cover-letter path (original layout, now marker-safe) ─────────────

def add_docx_text_block(doc, line):
    stripped = line.strip()
    if not stripped:
        doc.add_paragraph()
        return

    if set(stripped) <= {"=", "_", "-"} and len(stripped) >= 8:
        return

    if stripped.startswith(("•", "* ", "- ")):
        bullet_text = re.sub(r"^[•*-]\s+", "", stripped, count=1).strip()
        para = doc.add_paragraph(style="List Bullet")
        para.paragraph_format.left_indent = Inches(0.25)
        add_runs_with_markup(para, bullet_text)
        return

    plain = stripped.replace("**", "")
    if plain.isupper() and len(plain) <= 80:
        para = doc.add_paragraph()
        para.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
        run = para.add_run(plain)
        run.bold = True
        run.font.size = Pt(12)
        return

    if " — " in stripped and not stripped.endswith(":"):
        left, right = stripped.split(" — ", 1)
        # "Company — Title" header: left half is plain text. If the left half
        # carries ** markers it is body prose (a sentence with a dash), not a
        # header — fall through so the markers become real bold runs.
        if "**" not in left:
            para = doc.add_paragraph()
            para.paragraph_format.space_after = Pt(3)
            run1 = para.add_run(left + " — ")
            run1.bold = True
            add_runs_with_markup(para, right)
            return

    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(4)
    add_runs_with_markup(para, stripped)


def _is_effectively_blank(line):
    """Blank line OR a ---/===/___ separator — both collapse to one gap in the docx."""
    stripped = line.strip()
    if not stripped:
        return True
    return set(stripped) <= {"=", "_", "-"} and len(stripped) >= 8


# "Company — Title" header line, as used for each Professional Experience entry.
_COMPANY_HEADER_RE = re.compile(r"^\S.*\s—\s\S")

# "Mon YYYY – Mon YYYY" / "Mon YYYY – Present" date-range line that follows
# each company header line.
_DATE_RANGE_RE = re.compile(
    r"^[A-Za-z]{3,9}\.?\s+\d{4}\s*[–—-]\s*(?:[A-Za-z]{3,9}\.?\s+\d{4}|Present)$"
)


def _render_plain_text(doc, text):
    # Pass 1: a section boundary is "blank line, separator, blank line" in the
    # .txt — each of those independently would add its own empty paragraph,
    # doubling the visual gap in Word. Collapse any run of blank/separator
    # lines into a single blank paragraph between real content lines.
    collapsed_lines = []
    blank_pending = False
    for line in text.splitlines():
        if _is_effectively_blank(line):
            blank_pending = True
            continue
        if blank_pending and collapsed_lines:
            collapsed_lines.append("")
        blank_pending = False
        collapsed_lines.append(line)

    # Pass 2: within a Professional Experience entry, drop the blank-line gap
    # between the "Company — Title" header and its date-range line right
    # below it, so those two rows sit adjacent with no spacing.
    normalized_lines = []
    i = 0
    while i < len(collapsed_lines):
        line = collapsed_lines[i]
        if (
            line == ""
            and 0 < i < len(collapsed_lines) - 1
            and _COMPANY_HEADER_RE.match(collapsed_lines[i - 1].strip())
            and _DATE_RANGE_RE.match(collapsed_lines[i + 1].strip())
        ):
            i += 1
            continue
        normalized_lines.append(line)
        i += 1

    for line in normalized_lines:
        add_docx_text_block(doc, line)


# ── Markdown path (ScoreCards) ────────────────────────────────────────────────

_HR_RE = re.compile(r"^\s*([-*_=])(\s*\1){2,}\s*$")
_HEADING_RE = re.compile(r"^\s*(#{1,6})\s+(.*?)\s*#*\s*$")
_BULLET_RE = re.compile(r"^(\s*)[-*•+]\s+(.*)$")
_NUMBERED_RE = re.compile(r"^(\s*)(\d+)[.)]\s+(.*)$")
_QUOTE_RE = re.compile(r"^\s*>\s?(.*)$")
_TABLE_SEP_CELL_RE = re.compile(r"^:?-{2,}:?$")

_HEADING_STYLE = {   # level -> (size pt, space_before, space_after, navy colour?)
    1: (17, 0, 6, True),
    2: (13.5, 12, 4, True),
    3: (12, 9, 3, False),
    4: (11, 7, 2, False),
}


def _split_table_row(row):
    return [c.strip() for c in row.strip().strip("|").split("|")]


def _shade_cell(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def _add_table(doc, block):
    rows = [_split_table_row(r) for r in block]
    # drop empty rows and the |---|---| separator row
    rows = [r for r in rows if any(r) and not all(_TABLE_SEP_CELL_RE.match(c) for c in r if c)]
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    for ri, row in enumerate(rows):
        for ci in range(ncols):
            cell = table.cell(ri, ci)
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(1)
            para.paragraph_format.space_before = Pt(1)
            add_runs_with_markup(para, row[ci] if ci < len(row) else "", bold=(ri == 0), size=10)
            if ri == 0:
                _shade_cell(cell, "E7ECF3")
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)


def render_markdown(doc, text):
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        raw = lines[i]
        s = raw.strip()
        if not s or _HR_RE.match(raw):
            i += 1
            continue

        m = _HEADING_RE.match(raw)
        if m:
            level = min(len(m.group(1)), 4)
            size, before, after, navy = _HEADING_STYLE[level]
            para = doc.add_paragraph()
            para.paragraph_format.space_before = Pt(before)
            para.paragraph_format.space_after = Pt(after)
            para.paragraph_format.keep_with_next = True
            add_runs_with_markup(para, m.group(2), bold=True, size=size)
            if navy:
                for run in para.runs:
                    run.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
            i += 1
            continue

        if s.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            _add_table(doc, block)
            continue

        m = _BULLET_RE.match(raw)
        if m and not s.startswith("**"):
            level2 = len(m.group(1).replace("\t", "    ")) >= 2
            para = doc.add_paragraph(style="List Bullet 2" if level2 else "List Bullet")
            para.paragraph_format.space_after = Pt(2)
            add_runs_with_markup(para, m.group(2))
            i += 1
            continue

        m = _NUMBERED_RE.match(raw)
        if m:
            para = doc.add_paragraph()
            pf = para.paragraph_format
            pf.left_indent = Inches(0.3)
            pf.first_line_indent = Inches(-0.3)
            pf.space_after = Pt(2)
            pf.tab_stops.add_tab_stop(Inches(0.3))
            _emit_run(para, f"{m.group(2)}.\t")
            add_runs_with_markup(para, m.group(3))
            i += 1
            continue

        m = _QUOTE_RE.match(raw)
        if m:
            para = doc.add_paragraph()
            para.paragraph_format.left_indent = Inches(0.3)
            para.paragraph_format.space_after = Pt(3)
            add_runs_with_markup(para, m.group(1))
            for run in para.runs:
                run.italic = True
            i += 1
            continue

        para = doc.add_paragraph()
        para.paragraph_format.space_after = Pt(4)
        add_runs_with_markup(para, s)
        i += 1


def _looks_like_markdown(txt_path, text):
    if txt_path.parent.parent.name == "ScoreCard":
        return True
    return any(_HEADING_RE.match(l) for l in text.splitlines())


def convert_text_file_to_docx(txt_path):
    if not DOCX_AVAILABLE:
        raise RuntimeError("python-docx is not installed")

    txt_path = Path(txt_path)
    text = txt_path.read_text(encoding="utf-8")
    doc = Document()
    style_docx_document(doc)

    if _looks_like_markdown(txt_path, text):
        render_markdown(doc, text)
    else:
        _render_plain_text(doc, text)

    target_dir = txt_path.parent.parent / "docx"
    target_dir.mkdir(parents=True, exist_ok=True)
    out_path = target_dir / f"{txt_path.stem}.docx"
    try:
        doc.save(out_path)
        return out_path, False
    except PermissionError:
        fallback_name = f"{txt_path.stem}_{datetime.now().strftime('%H%M%S')}.docx"
        fallback_path = target_dir / fallback_name
        doc.save(fallback_path)
        return fallback_path, True


# ── Standalone rebuild CLI ────────────────────────────────────────────────────

def _main(argv):
    root = Path(__file__).resolve().parent.parent
    args = [a for a in argv if not a.startswith("-")]
    backup = "--no-backup" not in argv
    if not args:
        sys.exit(__doc__)
    target = Path(args[0])
    if not target.is_absolute():
        target = root / target
    if target.is_dir():
        txts = sorted(p for p in target.rglob("*.txt") if p.parent.name == "txt")
    elif target.suffix.lower() == ".txt" and target.exists():
        txts = [target]
    else:
        sys.exit(f"ERROR: need an existing .txt file or folder: {target}")
    if not txts:
        sys.exit(f"ERROR: no output .txt files found under {target}")

    stamp = datetime.now().strftime("%Y%m%d")
    failed = 0
    for txt in txts:
        existing = txt.parent.parent / "docx" / f"{txt.stem}.docx"
        if backup and existing.exists():
            try:
                rel = existing.relative_to(root)
            except ValueError:
                rel = Path(existing.name)
            bak = root / "archive" / f"docx_rebuild_{stamp}" / rel
            bak.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(existing, bak)
        try:
            out, fallback = convert_text_file_to_docx(txt)
            print(f"OK   {out}{'  (original locked, wrote fallback)' if fallback else ''}")
        except Exception as exc:
            failed += 1
            print(f"FAIL {txt}: {exc}")
    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    _main(sys.argv[1:])
